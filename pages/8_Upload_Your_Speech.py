import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
from textblob import TextBlob
from docx import Document
import PyPDF2
from fpdf import FPDF
import os

st.set_page_config(page_title="Upload Your Speech", layout="wide")

st.title("Upload Your Speech for Feedback")

st.markdown("""
Upload your script as `.txt`, `.docx`, or `.pdf`, or paste your content below.  
Click **"Run Analysis"** to generate feedback based on TED Talk benchmarks.
""")

@st.cache_data
def load_benchmarks():
    return pd.read_csv("data/processed/sentence_metrics.csv")

def extract_text_from_docx(file):
    try:
        doc = Document(file)
        return "\n".join([para.text for para in doc.paragraphs])
    except:
        return None

def extract_text_from_pdf(file):
    try:
        reader = PyPDF2.PdfReader(file)
        return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
    except:
        return None

def analyze_text(text):
    blob = TextBlob(text)
    sentences = blob.sentences
    sentiment = blob.sentiment

    rhetorical = [s for s in sentences if '?' in s and any(w in s.lower() for w in ['what', 'why', 'how'])]
    imperatives = [s for s in sentences if s.lower().startswith(('let', 'imagine', 'consider', 'try', 'think'))]

    clarity = 1.0 if len(sentences) == 0 else max(0, 1 - abs((len(blob.words) / len(sentences)) - 17) / 17)
    engagement = min(1.0, (len(rhetorical) + len(imperatives)) / max(1, len(sentences)))

    return {
        "sentence_count": len(sentences),
        "word_count": len(blob.words),
        "avg_sentence_length": len(blob.words) / len(sentences) if sentences else 0,
        "sentiment_polarity": sentiment.polarity,
        "sentiment_subjectivity": sentiment.subjectivity,
        "rhetorical_count": len(rhetorical),
        "imperative_count": len(imperatives),
        "rhetorical_samples": rhetorical[:3],
        "imperative_samples": imperatives[:3],
        "clarity_score": round(clarity * 100, 1),
        "engagement_score": round(engagement * 100, 1)
    }

# Input
uploaded_file = st.file_uploader("Upload a .txt, .docx, or .pdf file", type=["txt", "docx", "pdf"])
manual_text = st.text_area("Or paste your speech here:")

col1, col2 = st.columns([1, 1])
run_analysis = col1.button("Run Analysis")
reset = col2.button("Clear All")

if reset:
    st.session_state.clear()
    st.rerun()

user_text = None
if uploaded_file:
    file_ext = uploaded_file.name.lower().split(".")[-1]
    if file_ext == "txt":
        user_text = uploaded_file.read().decode("utf-8")
    elif file_ext == "docx":
        user_text = extract_text_from_docx(uploaded_file)
    elif file_ext == "pdf":
        user_text = extract_text_from_pdf(uploaded_file)
    else:
        st.error("Unsupported file type.")
elif manual_text.strip():
    user_text = manual_text.strip()

# Main analysis block
if user_text and run_analysis:
    st.markdown("### Speech Preview")
    st.text_area("Speech Text", user_text, height=200)

    result = analyze_text(user_text)
    df_benchmark = load_benchmarks()
    avg_benchmark = df_benchmark["avg_sentence_length"].mean()

    st.markdown("### Key Metrics")
    chart_data = pd.DataFrame({
        "Metric": ["Average Sentence Length", "Sentiment Polarity", "Sentiment Subjectivity"],
        "Score": [round(result["avg_sentence_length"], 2),
                  round(result["sentiment_polarity"], 2),
                  round(result["sentiment_subjectivity"], 2)]
    })

    fig, ax = plt.subplots()
    ax.barh(chart_data["Metric"], chart_data["Score"], color="#4682B4")
    ax.set_xlabel("Value")
    st.pyplot(fig)

    st.markdown("### TED Benchmark Comparison")
    st.write(f"TED Talk Average Sentence Length: **{round(avg_benchmark, 2)} words**")
    if result["avg_sentence_length"] > avg_benchmark:
        st.warning("Your sentences are longer than typical TED Talks. Consider breaking them down.")
    elif result["avg_sentence_length"] < avg_benchmark:
        st.info("Your sentences are shorter than average. This may improve clarity but limit depth.")
    else:
        st.success("Your sentence length aligns well with TED averages.")

    st.markdown("### Rhetorical Questions")
    st.write(f"Count: {result['rhetorical_count']}")
    for i, s in enumerate(result["rhetorical_samples"], 1):
        st.markdown(f"*{i}. {s}*")

    st.markdown("### Imperative Statements")
    st.write(f"Count: {result['imperative_count']}")
    for i, s in enumerate(result["imperative_samples"], 1):
        st.markdown(f"*{i}. {s}*")

    st.markdown("### Feedback Quality Score")
    fig2, ax2 = plt.subplots()
    ax2.barh(["Clarity", "Engagement"],
             [result["clarity_score"], result["engagement_score"]],
             color=["#2c7be5", "#00b894"])
    ax2.set_xlabel("Score (%)")
    ax2.set_xlim(0, 100)
    st.pyplot(fig2)

    # Downloadable CSV
    feedback_df = pd.DataFrame({
        "Metric": [
            "Word Count",
            "Sentence Count",
            "Average Sentence Length",
            "Sentiment Polarity",
            "Sentiment Subjectivity",
            "Rhetorical Questions",
            "Imperative Statements",
            "Clarity Score",
            "Engagement Score"
        ],
        "Value": [
            result["word_count"],
            result["sentence_count"],
            round(result["avg_sentence_length"], 2),
            round(result["sentiment_polarity"], 2),
            round(result["sentiment_subjectivity"], 2),
            result["rhetorical_count"],
            result["imperative_count"],
            result["clarity_score"],
            result["engagement_score"]
        ]
    })

    csv_path = "files/csv/speech_feedback_summary.csv"
    feedback_df.to_csv(csv_path, index=False)

    with open(csv_path, "rb") as f:
        st.download_button("Download Feedback as CSV", f, file_name=csv_path, mime="text/csv")

    # Generate PDF
    pdf_path = "files/pdf/speech_feedback_summary.pdf"

    class PDF(FPDF):
        def header(self):
            self.set_font("Arial", "B", 14)
            self.cell(0, 10, "SpeakScape Speech Analysis Report", ln=1, align="C")
            self.ln(5)

        def section_title(self, title):
            self.set_font("Arial", "B", 12)
            self.cell(0, 10, title, ln=1)
            self.set_font("Arial", "", 12)

        def metric_row(self, name, value):
            self.cell(80, 10, name, 0, 0)
            self.cell(40, 10, str(value), 0, 1)

    pdf = PDF()
    pdf.add_page()
    pdf.section_title("Key Metrics")
    for i in range(len(feedback_df)):
        pdf.metric_row(feedback_df["Metric"][i], feedback_df["Value"][i])

    pdf.ln(10)
    pdf.section_title("Recommendations")
    pdf.multi_cell(0, 10, "- Consider shortening long sentences to improve clarity.\n"
                         "- Increase rhetorical or imperative usage to boost engagement.\n"
                         "- Aim for balance in tone to match TED style.")

    pdf.output(pdf_path)

    with open(pdf_path, "rb") as f:
        st.download_button("Download Feedback as PDF", f, file_name=pdf_path, mime="application/pdf")

st.markdown("---")
st.caption("Built with ❤️ using Streamlit | Powered by TED Talks | Data from Kaggle")
